"""JSON / Excel / zip writers."""
import json
import os
import shutil
import subprocess
import tempfile
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

F = "Arial"
THIN = Side(style="thin", color="999999")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)


def write_json(summary, path):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(summary, f, indent=2, ensure_ascii=False)


def write_xlsx(summary, path):
    opt = summary["options"]
    ms = summary["members"]
    wb = Workbook()
    ws = wb.active
    ws.title = "Members"
    wa = wb.create_sheet("Assumptions")
    wp = wb.create_sheet("Pages")
    wn = wb.create_sheet("Notes")

    wa["A1"] = "Assumptions (edit the blue cells; lengths recalculate)"
    wa["A1"].font = Font(name=F, bold=True, size=12)
    rows = [("Paper mm per PDF unit", opt["mm_per_unit"], "1 PDF unit = 1 pt = 1/72 inch = 0.352778 mm"),
            ("Drawing scale denominator", opt["scale"], "Drawing scale 1:N, used for the real-world length"),
            ("Annotation offset (PDF units)", opt["offset"], "Blue line distance from member (reference only)")]
    for i, (a, v, c) in enumerate(rows, 3):
        wa.cell(i, 1, a).font = Font(name=F)
        x = wa.cell(i, 2, v)
        x.font = Font(name=F, color="0000FF")
        x.fill = PatternFill("solid", fgColor="FFFF99")
        wa.cell(i, 3, c).font = Font(name=F, italic=True)
    wa.column_dimensions["A"].width = 32
    wa.column_dimensions["B"].width = 12
    wa.column_dimensions["C"].width = 70

    ws["A1"] = f"{summary['member']} steel members - {summary['input']}"
    ws["A1"].font = Font(name=F, bold=True, size=12)
    ws["A2"] = ("Coordinates and base length in PDF units (origin bottom-left, y up). Start = left end (lower end if vertical). "
                "Paper and real lengths are formulas using the Assumptions sheet.")
    ws["A2"].font = Font(name=F, italic=True, size=9)
    hdr = ["ID", "Label", "Page", "Orientation", "Angle (deg, y-up)", "X1", "Y1", "X2", "Y2", "Length (PDF units)",
           "Length on paper (mm)", "Real length @ scale (mm)", "Label-to-member dist", "Next-nearest member dist",
           "Label match score", "Label read by OCR", "Segments merged", "Annotation side",
           "Blue line X1", "Blue line Y1", "Blue line X2", "Blue line Y2", "Found by", "Line colour / style"]
    HR = 4
    for j, h in enumerate(hdr, 1):
        c = ws.cell(HR, j, h)
        c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        c.border = BOX
    ws.row_dimensions[HR].height = 40
    r0 = HR + 1
    for i, m in enumerate(ms):
        r = r0 + i
        vals = {1: m["id"], 2: m["label"], 3: m["page"], 4: m["orientation"], 6: m["start"]["x_pdf"], 7: m["start"]["y_pdf"],
                8: m["end"]["x_pdf"], 9: m["end"]["y_pdf"], 13: m["label_to_member_distance"],
                14: m["next_nearest_member_distance"], 15: m["label_match"]["score"], 16: m["label_match"]["ocr_status"],
                17: m["n_segments_merged"], 18: m["annotation"]["side"], 19: m["annotation"]["start"]["x_pdf"],
                20: m["annotation"]["start"]["y_pdf"], 21: m["annotation"]["end"]["x_pdf"], 22: m["annotation"]["end"]["y_pdf"],
                23: m.get("found_by", "bold"), 24: f'{m["annotation"].get("style", {}).get("color", "")} {m["annotation"].get("style", {}).get("dash", "")}'.strip()}
        for k, v in vals.items():
            ws.cell(r, k, v)
        ws.cell(r, 5, f"=MOD(DEGREES(ATAN2(H{r}-F{r},I{r}-G{r})),180)")
        ws.cell(r, 10, f"=SQRT((H{r}-F{r})^2+(I{r}-G{r})^2)")
        ws.cell(r, 11, f"=J{r}*Assumptions!$B$3")
        ws.cell(r, 12, f"=K{r}*Assumptions!$B$4")
        for j in range(1, len(hdr) + 1):
            c = ws.cell(r, j)
            c.font = Font(name=F, color="0000FF" if j in (6, 7, 8, 9) else "000000")
            c.border = BOX
            if j == 5:
                c.number_format = "0.0"
            elif j in (6, 7, 8, 9, 10, 13, 14, 19, 20, 21, 22):
                c.number_format = "#,##0.00"
            elif j == 11:
                c.number_format = "#,##0.0"
            elif j == 12:
                c.number_format = "#,##0"
            elif j == 15:
                c.number_format = "0.00"
    last = r0 + len(ms) - 1
    tr = last + 2
    ws.cell(tr, 1, "Members").font = Font(name=F, bold=True)
    ws.cell(tr, 2, f"=COUNTA(A{r0}:A{last})" if ms else 0).font = Font(name=F, bold=True)
    ws.cell(tr + 1, 1, "Totals").font = Font(name=F, bold=True)
    if ms:
        for col, fmt in ((10, "#,##0.00"), (11, "#,##0.0"), (12, "#,##0")):
            L = get_column_letter(col)
            c = ws.cell(tr + 1, col, f"=SUM({L}{r0}:{L}{last})")
            c.font = Font(name=F, bold=True)
            c.number_format = fmt
    for j, w in enumerate([13, 9, 7, 12, 11, 11, 11, 11, 11, 14, 14, 16, 13, 14, 11, 20, 10, 14, 11, 11, 11, 11, 11, 18], 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(r0, 2)

    ph = ["Page", "Size (PDF units)", "Bold line width (pt)", "Label candidates", "Members annotated", "Rejected", "Warnings"]
    for j, h in enumerate(ph, 1):
        c = wp.cell(1, j, h)
        c.font = Font(name=F, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.border = BOX
    for i, p in enumerate(summary["pages"], 2):
        vals = [p["page"], f"{p['width']:.0f} x {p['height']:.0f}", p.get("bold_width"), p["candidates"], p["members"],
                "; ".join(f"{r['reason']}" for r in p["rejected"])[:300], "; ".join(p["warnings"])]
        for j, v in enumerate(vals, 1):
            c = wp.cell(i, j, v)
            c.font = Font(name=F)
            c.border = BOX
    for j, w in enumerate([7, 18, 18, 16, 18, 60, 60], 1):
        wp.column_dimensions[get_column_letter(j)].width = w

    notes = ["Method",
             "Text in CAD-exported PDFs is usually outlined (no text layer), so labels are found with OCR plus glyph template matching.",
             "Each label is paired with the nearest bold (thickest black) line run that is parallel to the text. Collinear touching segments are merged into one member.",
             "Longer marks (e.g. SB4A when searching SB4) are rejected by re-reading each candidate with OCR.",
             "Annotation: 1.44 pt line, same length, offset to the right (vertical members) or below (horizontal); ID placed on the outer side. Solid blue = standard finder.",
             "Dotted lines in another colour = the standard finder found no member and a fallback did (column 'Found by'): colour = coloured/dashed line, width = other line width, leader = via leader line, wide = wider search. Check those members on the drawing.",
             "", "Units",
             "Base length is in PDF units. Paper length (mm) = units x 0.352778. Real length (mm) = paper length x drawing scale.",
             "IDs are numbered continuously through the whole document (page order, left to right, top to bottom)."]
    for i, t in enumerate(notes, 1):
        wn.cell(i, 1, t).font = Font(name=F, bold=t in ("Method", "Units"))
    wn.column_dimensions["A"].width = 150
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    recalc_with_libreoffice(path)


def recalc_with_libreoffice(path):
    """Optional: if LibreOffice is installed, store calculated values in the file so any viewer shows them."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return False
    try:
        with tempfile.TemporaryDirectory() as td:
            subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", td, path],
                           check=True, capture_output=True, timeout=180)
            out = os.path.join(td, os.path.basename(path))
            if os.path.exists(out):
                shutil.copyfile(out, path)
                return True
    except Exception:
        return False
    return False


def write_zip(paths, zip_path):
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as z:
        for p in paths:
            z.write(p, os.path.basename(p))
